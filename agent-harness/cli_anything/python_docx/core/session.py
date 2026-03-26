"""Core session primitives for the python-docx harness."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from cli_anything.python_docx.core.errors import SessionError
from cli_anything.python_docx.utils import python_docx_backend as backend


class DocxSession:
    PALETTE_MAIN = "05206E"
    PALETTE_SECONDARY = "357AE9"
    PALETTE_WHITE = "FFFFFF"
    FRONTPAGE_TEMPLATES = ("clean", "corporate", "academic")
    BIBLIOGRAPHY_HEADING = "Bibliography"
    BIBLIOGRAPHY_KEY_PATTERN = re.compile(r"^[A-Za-z0-9._:-]+$")
    BIBLIOGRAPHY_ENTRY_PATTERN = re.compile(
        r"^\[(?P<number>\d+)\]\s+(?P<key>[A-Za-z0-9._:-]+)\s*:\s*(?P<reference>.+)$"
    )

    def __init__(self) -> None:
        self._document: Any | None = None
        self._path: Path | None = None
        self._undo_stack: list[bytes] = []
        self._redo_stack: list[bytes] = []

    @property
    def path(self) -> Path | None:
        return self._path

    @property
    def has_document(self) -> bool:
        return self._document is not None

    def new_document(self, path: str | Path | None = None, title: str | None = None) -> None:
        self._document = backend.new_document()
        self._path = Path(path) if path is not None else None
        self._undo_stack.clear()
        self._redo_stack.clear()
        if title:
            self.set_core_property("title", title)
        if path is not None:
            self.save(path)

    def open_document(self, path: str | Path) -> None:
        target = Path(path)
        if not target.exists():
            raise SessionError(f"Document does not exist: {target}")
        self._document = backend.load_document(target)
        self._path = target
        self._undo_stack.clear()
        self._redo_stack.clear()

    def save(self, path: str | Path | None = None) -> Path:
        self._ensure_document()
        target = Path(path) if path is not None else self._path
        if target is None:
            raise SessionError("No target path set. Pass a path to save().")
        target.parent.mkdir(parents=True, exist_ok=True)
        backend.save_document(self._document, target)
        self._path = target
        return target

    def summary(self) -> dict[str, Any]:
        doc = self._ensure_document()
        paragraphs = doc.paragraphs
        tables = doc.tables
        heading_count = 0
        for para in paragraphs:
            style = getattr(para, "style", None)
            style_name = getattr(style, "name", "")
            if isinstance(style_name, str) and style_name.startswith("Heading"):
                heading_count += 1

        return {
            "path": str(self._path) if self._path is not None else None,
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "heading_count": heading_count,
            "undo_depth": len(self._undo_stack),
            "redo_depth": len(self._redo_stack),
        }

    def list_paragraphs(self, limit: int | None = None) -> list[dict[str, Any]]:
        doc = self._ensure_document()
        records: list[dict[str, Any]] = []
        for index, para in enumerate(doc.paragraphs):
            style = getattr(para, "style", None)
            style_name = getattr(style, "name", None)
            records.append(
                {
                    "index": index,
                    "text": para.text,
                    "style": style_name,
                }
            )
        if limit is None:
            return records
        return records[: max(limit, 0)]

    def add_paragraph(self, text: str, style: str | None = None) -> None:
        doc = self._ensure_document()
        self._checkpoint()
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)

    def add_heading(self, text: str, level: int = 1) -> None:
        doc = self._ensure_document()
        self._checkpoint()
        paragraph = doc.add_heading(text, level=level)
        self._apply_paragraph_color(paragraph, self.PALETTE_MAIN)

    def list_frontpage_templates(self) -> list[str]:
        return list(self.FRONTPAGE_TEMPLATES)

    def add_frontpage(
        self,
        template: str,
        title: str,
        subtitle: str | None = None,
        author: str | None = None,
        organization: str | None = None,
        date_text: str | None = None,
        include_page_break: bool = True,
        set_core_title: bool = True,
    ) -> dict[str, Any]:
        doc = self._ensure_document()
        self._checkpoint()

        normalized_template = template.strip().lower()
        if normalized_template not in self.FRONTPAGE_TEMPLATES:
            raise SessionError(
                f"Unsupported frontpage template: {template}. "
                f"Available: {', '.join(self.FRONTPAGE_TEMPLATES)}"
            )
        if not title.strip():
            raise SessionError("Frontpage title cannot be empty.")

        body = doc._body._element
        existing_blocks = [child for child in body if child.tag != qn("w:sectPr")]
        old_first_block = existing_blocks[0] if existing_blocks else None

        effective_date_text = date_text or dt.date.today().isoformat()
        new_elements: list[Any] = []

        if normalized_template == "clean":
            self._append_frontpage_line(new_elements, doc, title, size=30, bold=True, space_after=18)
            if subtitle:
                self._append_frontpage_line(
                    new_elements,
                    doc,
                    subtitle,
                    size=16,
                    italic=True,
                    space_after=24,
                    color=self.PALETTE_SECONDARY,
                )
            if author:
                self._append_frontpage_line(new_elements, doc, author, size=12, space_after=6)
            if organization:
                self._append_frontpage_line(new_elements, doc, organization, size=12, space_after=6)
            self._append_frontpage_line(
                new_elements,
                doc,
                effective_date_text,
                size=11,
                space_after=0,
                color=self.PALETTE_SECONDARY,
            )
        elif normalized_template == "corporate":
            if organization:
                self._append_frontpage_line(
                    new_elements,
                    doc,
                    organization.upper(),
                    size=12,
                    bold=True,
                    space_after=24,
                    color=self.PALETTE_WHITE,
                    background_color=self.PALETTE_MAIN,
                )
            self._append_frontpage_line(
                new_elements,
                doc,
                title,
                size=28,
                bold=True,
                space_after=12,
                color=self.PALETTE_WHITE,
                background_color=self.PALETTE_MAIN,
            )
            if subtitle:
                self._append_frontpage_line(
                    new_elements,
                    doc,
                    subtitle,
                    size=14,
                    space_after=24,
                    color=self.PALETTE_WHITE,
                    background_color=self.PALETTE_MAIN,
                )
            if author:
                self._append_frontpage_line(
                    new_elements,
                    doc,
                    f"Prepared by {author}",
                    size=12,
                    space_after=6,
                    color=self.PALETTE_WHITE,
                    background_color=self.PALETTE_MAIN,
                )
            self._append_frontpage_line(
                new_elements,
                doc,
                effective_date_text,
                size=11,
                space_after=0,
                color=self.PALETTE_WHITE,
                background_color=self.PALETTE_MAIN,
            )
        else:
            self._append_frontpage_line(new_elements, doc, title, size=26, bold=True, space_after=12)
            if subtitle:
                self._append_frontpage_line(
                    new_elements,
                    doc,
                    subtitle,
                    size=14,
                    italic=True,
                    space_after=18,
                    color=self.PALETTE_SECONDARY,
                )
            if author:
                self._append_frontpage_line(new_elements, doc, f"Author: {author}", size=12, space_after=6)
            if organization:
                self._append_frontpage_line(new_elements, doc, f"Institution: {organization}", size=12, space_after=6)
            self._append_frontpage_line(
                new_elements,
                doc,
                effective_date_text,
                size=11,
                space_after=0,
                color=self.PALETTE_SECONDARY,
            )

        if include_page_break:
            break_para = doc.add_paragraph()
            break_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break_para.add_run().add_break(WD_BREAK.PAGE)
            new_elements.append(break_para._p)

        if old_first_block is not None:
            for element in new_elements:
                body.remove(element)
            insert_at = body.index(old_first_block)
            for offset, element in enumerate(new_elements):
                body.insert(insert_at + offset, element)

        if set_core_title:
            doc.core_properties.title = title

        return {
            "template": normalized_template,
            "title": title,
            "subtitle": subtitle,
            "author": author,
            "organization": organization,
            "date_text": effective_date_text,
            "page_break": include_page_break,
            "set_core_title": set_core_title,
        }

    def add_table(
        self,
        rows: int | None = None,
        cols: int | None = None,
        headers: list[str] | None = None,
        records: list[list[str]] | None = None,
        header_bold: bool = False,
        header_bg_color: str | None = None,
        row_lines: bool = False,
        column_lines: bool = False,
        outer_border: bool = False,
        line_style: str = "single",
        line_size: int = 8,
        line_color: str = "main",
    ) -> dict[str, int]:
        doc = self._ensure_document()
        self._checkpoint()
        headers = headers or []
        records = records or []

        has_structured_data = bool(headers or records)
        if not has_structured_data:
            if rows is None or cols is None:
                raise SessionError("rows and cols are required when no headers/records are provided.")
            table = doc.add_table(rows=rows, cols=cols)
            if row_lines or column_lines or outer_border:
                self._apply_table_borders(
                    table=table,
                    row_lines=row_lines,
                    column_lines=column_lines,
                    outer_border=outer_border,
                    line_style=line_style,
                    line_size=line_size,
                    line_color=line_color,
                )
            return {"rows": rows, "cols": cols}

        if cols is not None and cols < 1:
            raise SessionError("cols must be >= 1 when provided.")
        if rows is not None and rows < 1:
            raise SessionError("rows must be >= 1 when provided.")

        inferred_cols = cols or 0
        if headers:
            inferred_cols = max(inferred_cols, len(headers))
        for row in records:
            inferred_cols = max(inferred_cols, len(row))

        if inferred_cols < 1:
            raise SessionError("Unable to infer table column count from headers/records.")

        total_rows = (1 if headers else 0) + len(records)
        if rows is not None:
            total_rows = max(total_rows, rows)

        table = doc.add_table(rows=total_rows, cols=inferred_cols)

        row_index = 0
        if headers:
            normalized_header_bg_color = (
                self._resolve_palette_color(header_bg_color, "header_bg_color")
                if header_bg_color
                else self.PALETTE_MAIN
            )
            for col_index, value in enumerate(headers):
                cell = table.rows[0].cells[col_index]
                self._set_cell_text(cell, value, bold=header_bold, color=self.PALETTE_WHITE)
                self._set_cell_background(cell, normalized_header_bg_color)
            row_index = 1

        for record in records:
            cells = table.rows[row_index].cells
            for col_index, value in enumerate(record):
                cells[col_index].text = value
            row_index += 1

        if row_lines or column_lines or outer_border:
            self._apply_table_borders(
                table=table,
                row_lines=row_lines,
                column_lines=column_lines,
                outer_border=outer_border,
                line_style=line_style,
                line_size=line_size,
                line_color=line_color,
            )

        return {"rows": total_rows, "cols": inferred_cols}

    def add_bibliography_entry(
        self,
        key: str,
        reference: str,
        url: str | None = None,
    ) -> dict[str, Any]:
        doc = self._ensure_document()
        normalized_key = self._normalize_bibliography_key(key)
        normalized_reference = reference.strip()
        if not normalized_reference:
            raise SessionError("Bibliography reference cannot be empty.")
        normalized_url = url.strip() if url else None
        if normalized_url == "":
            normalized_url = None

        existing_entries = self._bibliography_entries(doc)
        existing_keys = {entry["key"].lower() for entry in existing_entries}
        if normalized_key.lower() in existing_keys:
            raise SessionError(f"Bibliography key already exists: {normalized_key}")

        self._checkpoint()

        _, insert_before = self._ensure_bibliography_heading(doc)
        existing_entries = self._bibliography_entries(doc)
        citation_number = max((entry["number"] for entry in existing_entries), default=0) + 1
        entry_text = f"[{citation_number}] {normalized_key}: {normalized_reference}"
        if normalized_url:
            entry_text = f"{entry_text} (URL: {normalized_url})"

        if insert_before is None:
            paragraph = doc.add_paragraph(entry_text)
        else:
            paragraph = insert_before.insert_paragraph_before(entry_text)

        return {
            "number": citation_number,
            "key": normalized_key,
            "reference": normalized_reference,
            "url": normalized_url,
            "text": paragraph.text,
        }

    def list_bibliography(self) -> list[dict[str, Any]]:
        doc = self._ensure_document()
        entries = self._bibliography_entries(doc)
        return [
            {
                "number": entry["number"],
                "key": entry["key"],
                "reference": entry["reference"],
            }
            for entry in entries
        ]

    def add_citation(
        self,
        text: str,
        source_keys: list[str],
        style: str | None = None,
    ) -> dict[str, Any]:
        doc = self._ensure_document()
        normalized_text = text.strip()
        if not normalized_text:
            raise SessionError("Citation text cannot be empty.")

        numbers = self._citation_numbers_for_keys(doc, source_keys)
        marker = self._format_citation_marker(numbers)
        paragraph_text = f"{normalized_text} {marker}"

        self._checkpoint()
        if style:
            paragraph = doc.add_paragraph(paragraph_text, style=style)
        else:
            paragraph = doc.add_paragraph(paragraph_text)

        return {
            "text": paragraph.text,
            "source_keys": [self._normalize_bibliography_key(key) for key in source_keys],
            "citation_numbers": numbers,
            "citation_marker": marker,
            "style": style,
        }

    def cite_paragraph(self, paragraph_index: int, source_keys: list[str]) -> dict[str, Any]:
        doc = self._ensure_document()
        paragraphs = doc.paragraphs
        if paragraph_index < 0 or paragraph_index >= len(paragraphs):
            raise SessionError(f"paragraph_index out of range: {paragraph_index}")

        numbers = self._citation_numbers_for_keys(doc, source_keys)
        marker = self._format_citation_marker(numbers)
        paragraph = paragraphs[paragraph_index]
        prior_text = paragraph.text
        updated_text = prior_text.rstrip()
        if not updated_text.endswith(marker):
            updated_text = f"{updated_text} {marker}".strip()

        self._checkpoint()
        paragraph.text = updated_text

        return {
            "paragraph_index": paragraph_index,
            "text_before": prior_text,
            "text_after": paragraph.text,
            "source_keys": [self._normalize_bibliography_key(key) for key in source_keys],
            "citation_numbers": numbers,
            "citation_marker": marker,
        }

    def _append_frontpage_line(
        self,
        elements: list[Any],
        doc: Any,
        text: str,
        *,
        size: int,
        bold: bool = False,
        italic: bool = False,
        space_after: int = 0,
        color: str | None = None,
        background_color: str | None = None,
    ) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if space_after > 0:
            paragraph.paragraph_format.space_after = Pt(space_after)
        if background_color:
            self._set_paragraph_background(paragraph, background_color)
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        self._set_run_color(run, color or self.PALETTE_MAIN)
        elements.append(paragraph._p)

    def _set_paragraph_background(self, paragraph: Any, fill_color: str) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        shd = p_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            p_pr.append(shd)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_color)

    def _set_cell_text(self, cell: Any, value: str, bold: bool = False, color: str | None = None) -> None:
        cell.text = value
        for paragraph in cell.paragraphs:
            if not paragraph.runs:
                paragraph.add_run("")
            for run in paragraph.runs:
                run.bold = bold
                self._set_run_color(run, color or self.PALETTE_MAIN)

    def _set_cell_background(self, cell: Any, fill_color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_color)

    def _normalize_hex_color(self, value: str, field_name: str) -> str:
        normalized = value.strip().lstrip("#").upper()
        if len(normalized) != 6 or any(ch not in "0123456789ABCDEF" for ch in normalized):
            raise SessionError(f"{field_name} must be a 6-digit hex value like 'D9E1F2'.")
        return normalized

    def _resolve_palette_color(self, value: str, field_name: str) -> str:
        normalized = value.strip().lower().lstrip("#")
        if normalized in {"main", self.PALETTE_MAIN.lower()}:
            return self.PALETTE_MAIN
        if normalized in {"secondary", self.PALETTE_SECONDARY.lower()}:
            return self.PALETTE_SECONDARY
        raise SessionError(
            f"{field_name} must be one of: main, secondary, {self.PALETTE_MAIN}, {self.PALETTE_SECONDARY}."
        )

    def _set_run_color(self, run: Any, color_hex: str) -> None:
        run.font.color.rgb = RGBColor.from_string(color_hex)

    def _apply_paragraph_color(self, paragraph: Any, color_hex: str) -> None:
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            self._set_run_color(run, color_hex)

    def _apply_table_borders(
        self,
        table: Any,
        row_lines: bool,
        column_lines: bool,
        outer_border: bool,
        line_style: str,
        line_size: int,
        line_color: str,
    ) -> None:
        color = self._resolve_palette_color(line_color, "line_color")

        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            tbl.insert(0, tbl_pr)

        tbl_borders = tbl_pr.find(qn("w:tblBorders"))
        if tbl_borders is None:
            tbl_borders = OxmlElement("w:tblBorders")
            tbl_pr.append(tbl_borders)

        border_edges: list[str] = []
        if outer_border:
            border_edges.extend(["top", "left", "bottom", "right"])
        if row_lines:
            border_edges.append("insideH")
        if column_lines:
            border_edges.append("insideV")

        for edge in border_edges:
            edge_elm = tbl_borders.find(qn(f"w:{edge}"))
            if edge_elm is None:
                edge_elm = OxmlElement(f"w:{edge}")
                tbl_borders.append(edge_elm)
            edge_elm.set(qn("w:val"), line_style)
            edge_elm.set(qn("w:sz"), str(line_size))
            edge_elm.set(qn("w:space"), "0")
            edge_elm.set(qn("w:color"), color)

    def _normalize_bibliography_key(self, key: str) -> str:
        normalized = key.strip()
        if not normalized:
            raise SessionError("Bibliography key cannot be empty.")
        if not self.BIBLIOGRAPHY_KEY_PATTERN.fullmatch(normalized):
            raise SessionError(
                "Bibliography key must only contain letters, numbers, dot, underscore, colon, or dash."
            )
        return normalized

    def _find_bibliography_heading(self, doc: Any) -> tuple[Any | None, int | None]:
        for index, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip().lower() == self.BIBLIOGRAPHY_HEADING.lower():
                return paragraph, index
        return None, None

    def _next_heading_after_index(self, doc: Any, start_index: int) -> int | None:
        for paragraph_index, paragraph in enumerate(doc.paragraphs[start_index + 1 :], start=start_index + 1):
            style = getattr(paragraph, "style", None)
            style_name = getattr(style, "name", "")
            if isinstance(style_name, str) and style_name.startswith("Heading"):
                return paragraph_index
        return None

    def _ensure_bibliography_heading(self, doc: Any) -> tuple[Any, Any | None]:
        heading, heading_index = self._find_bibliography_heading(doc)
        if heading is None or heading_index is None:
            heading = doc.add_heading(self.BIBLIOGRAPHY_HEADING, level=1)
            self._apply_paragraph_color(heading, self.PALETTE_MAIN)
            return heading, None
        next_heading_index = self._next_heading_after_index(doc, heading_index)
        if next_heading_index is None:
            return heading, None
        return heading, doc.paragraphs[next_heading_index]

    def _bibliography_entries(self, doc: Any) -> list[dict[str, Any]]:
        _, heading_index = self._find_bibliography_heading(doc)
        if heading_index is None:
            return []

        next_heading_index = self._next_heading_after_index(doc, heading_index)
        if next_heading_index is None:
            end_index = len(doc.paragraphs)
        else:
            end_index = next_heading_index

        entries: list[dict[str, Any]] = []
        for paragraph_index in range(heading_index + 1, end_index):
            text = doc.paragraphs[paragraph_index].text.strip()
            if not text:
                continue
            match = self.BIBLIOGRAPHY_ENTRY_PATTERN.match(text)
            if not match:
                continue
            entries.append(
                {
                    "number": int(match.group("number")),
                    "key": match.group("key"),
                    "reference": match.group("reference").strip(),
                    "paragraph_index": paragraph_index,
                }
            )
        entries.sort(key=lambda entry: entry["number"])
        return entries

    def _citation_numbers_for_keys(self, doc: Any, source_keys: list[str]) -> list[int]:
        if not source_keys:
            raise SessionError("At least one source key is required.")

        entries = self._bibliography_entries(doc)
        if not entries:
            raise SessionError("No bibliography entries found. Add bibliography entries before citing.")

        key_to_number = {entry["key"].lower(): entry["number"] for entry in entries}
        citation_numbers: list[int] = []
        missing_keys: list[str] = []
        for key in source_keys:
            normalized_key = self._normalize_bibliography_key(key)
            number = key_to_number.get(normalized_key.lower())
            if number is None:
                missing_keys.append(normalized_key)
                continue
            if number not in citation_numbers:
                citation_numbers.append(number)

        if missing_keys:
            missing = ", ".join(missing_keys)
            raise SessionError(f"Unknown bibliography key(s): {missing}")

        return citation_numbers

    def _format_citation_marker(self, citation_numbers: list[int]) -> str:
        if len(citation_numbers) == 1:
            return f"[{citation_numbers[0]}]"
        joined = ", ".join(str(number) for number in citation_numbers)
        return f"[{joined}]"

    def set_core_property(self, key: str, value: str) -> None:
        doc = self._ensure_document()
        core = doc.core_properties
        if not hasattr(core, key):
            raise SessionError(f"Unsupported core property: {key}")
        self._checkpoint()
        setattr(core, key, value)

    def undo(self) -> None:
        self._ensure_document()
        if not self._undo_stack:
            raise SessionError("Nothing to undo.")
        current = backend.serialize_document(self._document)
        snapshot = self._undo_stack.pop()
        self._redo_stack.append(current)
        self._document = backend.load_document_from_bytes(snapshot)

    def redo(self) -> None:
        self._ensure_document()
        if not self._redo_stack:
            raise SessionError("Nothing to redo.")
        current = backend.serialize_document(self._document)
        snapshot = self._redo_stack.pop()
        self._undo_stack.append(current)
        self._document = backend.load_document_from_bytes(snapshot)

    def _checkpoint(self) -> None:
        self._ensure_document()
        snapshot = backend.serialize_document(self._document)
        self._undo_stack.append(snapshot)
        self._redo_stack.clear()

    def _ensure_document(self):
        if self._document is None:
            raise SessionError("No active document. Use new/open first.")
        return self._document

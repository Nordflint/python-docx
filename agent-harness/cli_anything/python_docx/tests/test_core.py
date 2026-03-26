from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from cli_anything.python_docx.core import DocxSession, SessionError


def it_handles_session_mutation_and_undo_redo() -> None:
    session = DocxSession()
    session.new_document()

    session.add_paragraph("first")
    assert session.summary()["paragraph_count"] == 1

    session.undo()
    assert session.summary()["paragraph_count"] == 0

    session.redo()
    assert session.summary()["paragraph_count"] == 1


def it_saves_opens_and_writes_core_property(tmp_path: Path) -> None:
    target = tmp_path / "report.docx"
    session = DocxSession()
    session.new_document(path=target, title="Draft")

    session.add_heading("Plan", level=2)
    session.add_table(rows=2, cols=2)
    session.set_core_property("author", "CLI Harness")
    session.save()

    reopened = DocxSession()
    reopened.open_document(target)

    summary = reopened.summary()
    assert summary["paragraph_count"] == 1
    assert summary["table_count"] == 1


def it_raises_when_undo_has_no_history() -> None:
    session = DocxSession()
    session.new_document()

    with pytest.raises(SessionError):
        session.undo()


def it_adds_structured_table_rows(tmp_path: Path) -> None:
    target = tmp_path / "records.docx"
    session = DocxSession()
    session.new_document(path=target)

    session.add_table(
        headers=["Qty", "Id", "Desc"],
        records=[
            ["3", "101", "Spam"],
            ["7", "422", "Eggs"],
            ["4", "631", "Spam, spam, eggs, and spam"],
        ],
        header_bold=True,
        header_bg_color="D9E1F2",
    )
    session.save()

    doc = Document(str(target))
    table = doc.tables[0]
    assert len(table.rows) == 4
    assert len(table.columns) == 3
    assert table.cell(0, 0).text == "Qty"
    assert table.cell(0, 1).text == "Id"
    assert table.cell(0, 2).text == "Desc"
    assert table.cell(1, 0).text == "3"
    assert table.cell(3, 2).text == "Spam, spam, eggs, and spam"
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    shd = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == "D9E1F2"


def it_applies_table_line_formatting(tmp_path: Path) -> None:
    target = tmp_path / "bordered.docx"
    session = DocxSession()
    session.new_document(path=target)
    session.add_table(
        rows=2,
        cols=2,
        row_lines=True,
        column_lines=True,
        outer_border=True,
        line_style="single",
        line_size=8,
        line_color="000000",
    )
    session.save()

    doc = Document(str(target))
    table = doc.tables[0]
    tbl_borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert tbl_borders is not None

    inside_h = tbl_borders.find(qn("w:insideH"))
    inside_v = tbl_borders.find(qn("w:insideV"))
    top = tbl_borders.find(qn("w:top"))
    assert inside_h is not None
    assert inside_v is not None
    assert top is not None
    assert inside_h.get(qn("w:val")) == "single"
    assert inside_v.get(qn("w:val")) == "single"
    assert inside_h.get(qn("w:color")) == "000000"

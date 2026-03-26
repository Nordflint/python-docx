from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _cli_cmd() -> list[str]:
    exe = shutil.which("cli-anything-python-docx")
    if exe:
        return [exe]
    return [sys.executable, "-m", "cli_anything.python_docx"]


def _run(args: list[str], input_text: str | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        _cli_cmd() + args,
        check=True,
        capture_output=True,
        text=True,
        input=input_text,
    )


def _json_lines(output: str) -> list[dict]:
    rows = []
    for line in output.splitlines():
        line = line.strip()
        start = line.find("{")
        if start < 0:
            continue
        candidate = line[start:]
        if not candidate.endswith("}"):
            continue
        rows.append(json.loads(candidate))
    return rows


def it_runs_one_shot_json_workflow(tmp_path: Path) -> None:
    doc_path = tmp_path / "workflow.docx"

    created = _run(["--json", "new", str(doc_path), "--title", "Workflow"])
    created_rows = _json_lines(created.stdout)
    assert created_rows and created_rows[-1]["ok"] is True

    _run(["add-paragraph", "Hello from e2e", "--doc", str(doc_path)])
    _run(["add-heading", "Milestone", "--level", "2", "--doc", str(doc_path)])

    summary = _run(["--json", "summary", "--doc", str(doc_path)])
    summary_rows = _json_lines(summary.stdout)
    assert summary_rows

    payload = summary_rows[-1]
    assert payload["summary"]["paragraph_count"] == 2
    assert payload["summary"]["heading_count"] == 1


def it_runs_default_repl_mode_with_undo() -> None:
    repl_script = "\n".join(
        [
            "new",
            "add-paragraph \"Temp line\"",
            "undo",
            "summary",
            "exit",
            "",
        ]
    )

    completed = _run(["--json"], input_text=repl_script)
    rows = _json_lines(completed.stdout)

    summary_rows = [row for row in rows if row.get("action") == "summary"]
    assert summary_rows
    assert summary_rows[-1]["summary"]["paragraph_count"] == 0


def it_adds_table_records_from_cli(tmp_path: Path) -> None:
    doc_path = tmp_path / "table-records.docx"
    _run(["new", str(doc_path)])
    _run(
        [
            "add-table",
            "--doc",
            str(doc_path),
            "--header",
            "Qty",
            "--header",
            "Id",
            "--header",
            "Desc",
            "--header-bold",
            "--header-bg-color",
            "D9E1F2",
            "--record",
            "3|101|Spam",
            "--record",
            "7|422|Eggs",
            "--record",
            "4|631|Spam, spam, eggs, and spam",
        ]
    )

    doc = Document(str(doc_path))
    table = doc.tables[0]
    assert len(table.rows) == 4
    assert len(table.columns) == 3
    assert table.cell(0, 0).text == "Qty"
    assert table.cell(1, 0).text == "3"
    assert table.cell(3, 2).text == "Spam, spam, eggs, and spam"
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    shd = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == "D9E1F2"


def it_adds_row_and_column_lines_from_cli(tmp_path: Path) -> None:
    doc_path = tmp_path / "table-lines.docx"
    _run(["new", str(doc_path)])
    _run(
        [
            "add-table",
            "--doc",
            str(doc_path),
            "--rows",
            "2",
            "--cols",
            "2",
            "--row-lines",
            "--column-lines",
            "--outer-border",
            "--line-style",
            "single",
            "--line-size",
            "8",
            "--line-color",
            "000000",
        ]
    )

    doc = Document(str(doc_path))
    table = doc.tables[0]
    tbl_borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert tbl_borders is not None
    assert tbl_borders.find(qn("w:insideH")) is not None
    assert tbl_borders.find(qn("w:insideV")) is not None

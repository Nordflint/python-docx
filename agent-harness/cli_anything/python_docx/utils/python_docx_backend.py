"""Backend wrapper over python-docx APIs."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document


def new_document() -> Any:
    return Document()


def load_document(path: str | Path) -> Any:
    return Document(str(path))


def save_document(document: Any, path: str | Path) -> None:
    document.save(str(path))


def serialize_document(document: Any) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def load_document_from_bytes(data: bytes) -> Any:
    stream = io.BytesIO(data)
    return Document(stream)